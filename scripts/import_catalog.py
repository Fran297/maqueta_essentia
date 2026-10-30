#!/usr/bin/env python3
"""Normalize the Altiva catalog workbook into the JSON used by both app layers."""
from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data/source/altiva_essentia_catalogo.xlsx"
OUTPUT = ROOT / "data/catalog.json"

FAMILY_IDS = {
    "Gourmand": "gourmand",
    "Amaderado / Oud": "oud",
    "Floral": "floral",
    "Ámbar Oriental": "amber",
    "Oriental Amaderado": "oriental-wood",
    "Cítrico / Fresco": "citrus",
    "Acuático / Marino": "aquatic",
    "Frutal": "fruity",
    "Oriental Especiada": "oriental-spicy",
    "Aromático Fougère": "aromatic-fougere",
    "Oriental": "oriental",
    "Aromático Acuático": "aromatic-aquatic",
    "Oriental Vainilla": "oriental-vanilla",
    "Floral Almizclado": "floral-musk",
}
GENDER_IDS = {"Unisex": "unisex", "Femenino": "fem", "Masculino": "masc"}


def text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def slug(value: str) -> str:
    folded = unicodedata.normalize("NFD", value.casefold())
    folded = "".join(ch for ch in folded if unicodedata.category(ch) != "Mn")
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", folded)).strip("-")


def whole_number(value: object) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    digits = re.sub(r"\D", "", str(value))
    return int(digits) if digits else None


def column_map(sheet, required_headers: tuple[str, ...]) -> tuple[list[str], list[tuple[object, ...]]]:
    rows = list(sheet.iter_rows(values_only=True))
    for position, row in enumerate(rows):
        headers = [text(cell) for cell in row]
        if all(header in headers for header in required_headers):
            return headers, rows[position + 1 :]
    raise ValueError(f"No se encontró la fila de encabezados en {sheet.title}.")


def parse_notes(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"[,/;]+", value) if part.strip()]


def resolve_photo(reference: str) -> str | None:
    if not reference:
        return None
    normalized = reference.replace("\\", "/").lstrip("/")
    if any(part in (".", "..") for part in Path(normalized).parts):
        return None
    candidate = ROOT / "public" / normalized
    return "/" + normalized if candidate.is_file() else None


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(f"No existe el archivo fuente: {SOURCE}")
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    if "Productos" not in workbook.sheetnames or "Casas" not in workbook.sheetnames:
        raise ValueError("El archivo debe incluir las hojas Productos y Casas.")

    product_headers, product_rows = column_map(workbook["Productos"], ("Casa", "Producto"))
    pindex = {name: index for index, name in enumerate(product_headers)}
    required = {
        "N°", "Casa", "Producto", "Destacado", "Familia olfativa", "Género",
        "Notas olfativas", "Presentación", "Tipo", "Volumen (ml)",
        "Año de lanzamiento", "Duración estimada", "Precio (ARS)",
        "Precio mostrado", "Foto (archivo)",
    }
    missing = required - set(pindex)
    if missing:
        raise ValueError(f"Faltan columnas requeridas en Productos: {sorted(missing)}")

    products: list[dict[str, object]] = []
    family_counts: Counter[str] = Counter()
    house_counts: Counter[str] = Counter()
    for row in product_rows:
        name = text(row[pindex["Producto"]])
        house_name = text(row[pindex["Casa"]])
        family_label = text(row[pindex["Familia olfativa"]])
        if not name or not house_name:
            continue
        if family_label not in FAMILY_IDS:
            raise ValueError(f"Familia sin mapeo explícito: {family_label!r}")
        gender_label = text(row[pindex["Género"]])
        if gender_label not in GENDER_IDS:
            raise ValueError(f"Género sin mapeo explícito: {gender_label!r}")

        price = whole_number(row[pindex["Precio (ARS)"]])
        raw_price_label = text(row[pindex["Precio mostrado"]])
        price_label = raw_price_label or (f"$ {price:,}" if price is not None else "Consultar")
        notes_raw = text(row[pindex["Notas olfativas"]])
        photo_ref = text(row[pindex["Foto (archivo)"]])
        house_id = slug(house_name)
        family_id = FAMILY_IDS[family_label]
        family_counts[family_label] += 1
        house_counts[house_name] += 1
        products.append({
            "id": f"ae-{whole_number(row[pindex['N°']]) or len(products) + 1:03d}",
            "sourceNumber": whole_number(row[pindex["N°"]]),
            "brand": house_id,
            "houseName": house_name,
            "name": name,
            "featured": text(row[pindex["Destacado"]]).casefold() in {"sí", "si", "yes", "true", "1"},
            "family": family_id,
            "familyLabel": family_label,
            "gender": GENDER_IDS[gender_label],
            "genderLabel": gender_label,
            "notesRaw": notes_raw,
            "notes": parse_notes(notes_raw),
            "presentation": text(row[pindex["Presentación"]]),
            "concentration": text(row[pindex["Tipo"]]),
            "volumeMl": whole_number(row[pindex["Volumen (ml)"]]),
            "launchYear": whole_number(row[pindex["Año de lanzamiento"]]),
            "launchYearSource": text(row[pindex["Año de lanzamiento"]]),
            "durationSource": text(row[pindex["Duración estimada"]]),
            "priceARS": price,
            "priceLabel": price_label,
            "inquiry": text(row[pindex.get("Consulta", -1)]) if "Consulta" in pindex else "",
            "photoRef": photo_ref,
            "imagePath": resolve_photo(photo_ref),
        })

    family_list = [
        {"id": FAMILY_IDS[label], "label": label, "count": count}
        for label, count in family_counts.most_common()
    ]
    house_headers, house_rows = column_map(workbook["Casas"], ("Casa", "Descripción"))
    hindex = {name: index for index, name in enumerate(house_headers)}
    houses: list[dict[str, object]] = []
    for row in house_rows:
        name = text(row[hindex["Casa"]])
        if not name or name.casefold().startswith(("total", "control:")):
            continue
        count_value = whole_number(row[hindex["Cantidad de fragancias"]]) if "Cantidad de fragancias" in hindex else None
        actual_count = house_counts[name]
        if count_value is not None and count_value != actual_count:
            raise ValueError(f"El conteo de {name} no coincide: Casas={count_value}, Productos={actual_count}")
        houses.append({
            "id": slug(name),
            "name": name,
            "tagline": text(row[hindex["Descripción"]]) if "Descripción" in hindex else "",
            "count": actual_count,
        })
    for name, count in house_counts.items():
        if not any(house["name"] == name for house in houses):
            houses.append({"id": slug(name), "name": name, "tagline": "", "count": count})

    if len(products) != 147 or len(houses) != 16 or len(family_list) != 14:
        raise ValueError(f"Conteos inesperados: productos={len(products)}, casas={len(houses)}, familias={len(family_list)}")
    if sum(product["priceARS"] is None for product in products) != 13:
        raise ValueError("Cambió el conteo de filas sin precio; revisar el Excel antes de regenerar.")
    if sum(product["featured"] for product in products) != 16:
        raise ValueError("Cambió el conteo de destacados; revisar el Excel antes de regenerar.")

    output = {
        "source": "data/source/altiva_essentia_catalogo.xlsx",
        "stats": {
            "products": len(products),
            "houses": len(houses),
            "families": len(family_list),
            "featured": sum(product["featured"] for product in products),
            "priceMissing": sum(product["priceARS"] is None for product in products),
            "durationMissingOrConsult": sum(not product["durationSource"] or product["durationSource"].casefold() == "consultar" for product in products),
            "volumeMissing": sum(product["volumeMl"] is None for product in products),
            "yearMissingOrConsult": sum(product["launchYear"] is None for product in products),
            "photoReferenceMissing": sum(not product["photoRef"] for product in products),
            "photoFilesAvailable": sum(bool(product["imagePath"]) for product in products),
        },
        "families": family_list,
        "houses": houses,
        "products": products,
    }
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(output["stats"], ensure_ascii=False, sort_keys=True))
    print(f"Wrote {len(products)} products to {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
